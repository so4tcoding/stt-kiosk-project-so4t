# 발표 슬라이드를 그림으로 그린 뒤 pptx에 넣는다.
# 숫자는 공개 조사 인용과 이 저장소에서 측정한 값만 사용한다.
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = Path("/workspace/presentation")
CHARTS = Path("/tmp/talk_charts")
SLIDES = Path("/tmp/talk_slides")
PAPER = Path("/workspace/paper")
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
OUT = ROOT / "kiosk_barrier_free_talk.pptx"

W, H = 1920, 1080
BG = (246, 240, 230)
INK = (44, 36, 22)
MUTED = (107, 94, 81)
ACCENT = (194, 81, 15)
NAVY = (31, 78, 121)
TEAL = (38, 128, 112)
CARD = (255, 252, 247)
LINE = (226, 214, 198)
GOLD = (184, 122, 36)
GREEN = (46, 125, 90)
RED = (166, 62, 48)

font_manager.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = "WenQuanYi Micro Hei"
plt.rcParams["axes.unicode_minus"] = False


def font(size):
    return ImageFont.truetype(FONT_PATH, size)


def wrap(draw, text, face, max_w):
    lines = []
    for para in str(text).split("\n"):
        if para == "":
            lines.append("")
            continue
        cur = ""
        for word in para.split(" "):
            trial = word if not cur else cur + " " + word
            if draw.textlength(trial, font=face) <= max_w:
                cur = trial
                continue
            if cur:
                lines.append(cur)
            if draw.textlength(word, font=face) <= max_w:
                cur = word
                continue
            buf = ""
            for ch in word:
                piece = buf + ch
                if draw.textlength(piece, font=face) <= max_w:
                    buf = piece
                else:
                    if buf:
                        lines.append(buf)
                    buf = ch
            cur = buf
        if cur:
            lines.append(cur)
    return lines


def draw_text(draw, xy, text, face, fill, max_w=None, leading=None):
    x, y = xy
    if max_w is None:
        draw.text((x, y), text, font=face, fill=fill)
        box = draw.textbbox((x, y), text, font=face)
        return box[3] + 4
    leading = leading or (face.size + 10)
    for line in wrap(draw, text, face, max_w):
        draw.text((x, y), line, font=face, fill=fill)
        y += leading
    return y


def rounded(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def new_slide(rail=True):
    image = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(image)
    if rail:
        draw.rectangle((0, 0, 14, H), fill=ACCENT)
    return image, draw


def footer(draw, page, total, label="초고령층을 위한 음성 키오스크"):
    draw.line((72, 1012, 1848, 1012), fill=LINE, width=2)
    draw.text((72, 1028), label, font=font(20), fill=MUTED)
    mark = f"{page}  /  {total}"
    face = font(20)
    width = draw.textlength(mark, font=face)
    draw.text((1848 - width, 1028), mark, font=face, fill=MUTED)


def header(draw, title, subtitle=""):
    y = draw_text(draw, (72, 32), title, font(40), INK, max_w=1760, leading=48)
    y += 8
    if subtitle:
        y = draw_text(draw, (72, y), subtitle, font(22), MUTED, max_w=1760, leading=30)
        y += 8
    draw.line((72, y, 1848, y), fill=LINE, width=2)
    return y + 16


def save_chart(name, fig):
    CHARTS.mkdir(parents=True, exist_ok=True)
    path = CHARTS / name
    fig.savefig(path, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_charts():
    # 서울디지털재단 2023 서울시민 디지털역량 실태조사, 2024-06-12 발표.
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    labels = ["55–64세", "65–74세", "75세 이상"]
    vals = [79.1, 50.4, 19.1]
    colors = ["#1F4E79", "#C47B2B", "#C2510F"]
    bars = ax.bar(labels, vals, color=colors, width=0.62)
    for bar, val in zip(bars, vals):
        ax.text(bar.get_x() + bar.get_width() / 2, val + 1.6, f"{val:.1f}%", ha="center", va="bottom", fontsize=14, color="#2C2416")
    ax.set_ylim(0, 100)
    ax.set_ylabel("키오스크 이용 경험 (%)", fontsize=12)
    ax.set_title("연령이 올라갈수록 키오스크 이용 경험이 줄어든다", fontsize=15, pad=12, color="#2C2416")
    ax.axhline(57.1, color="#6B5E51", ls="--", lw=1)
    ax.text(2.35, 60.2, "만 55세 이상 전체 57.1%", ha="right", fontsize=11, color="#6B5E51")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    age = save_chart("age_experience.png", fig)

    fig, ax = plt.subplots(figsize=(8.2, 4.2))
    reasons = ["뒷사람 눈치", "선택 사항이 어려움", "용어가 어려움"]
    rvals = [53.6, 46.3, 34.0]
    bars = ax.barh(reasons[::-1], rvals[::-1], color=["#C2510F", "#C47B2B", "#1F4E79"][::-1], height=0.62)
    for bar, val in zip(bars, rvals[::-1]):
        ax.text(val + 1.0, bar.get_y() + bar.get_height() / 2, f"{val:.1f}%", va="center", fontsize=13, color="#2C2416")
    ax.set_xlim(0, 78)
    ax.set_xlabel("어려움을 겪은 고령층 안의 비율 (%)  ·  복수 응답", fontsize=11)
    ax.set_title("키오스크가 어려웠던 이유", fontsize=15, pad=12, color="#2C2416")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    why = save_chart("difficulty_reasons.png", fig)

    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    names = ["고객이 한 말\n(겹침 제거)", "같은 뜻으로\n이어 주는 말", "사전 합계"]
    counts = [1274, 3595, 4869]
    colors = ["#1F4E79", "#268070", "#C2510F"]
    bars = ax.bar(names, counts, color=colors, width=0.62)
    for bar, val in zip(bars, counts):
        ax.text(bar.get_x() + bar.get_width() / 2, val + 80, f"{val:,}", ha="center", fontsize=14, color="#2C2416")
    ax.set_ylim(0, 5800)
    ax.set_ylabel("말의 수", fontsize=12)
    ax.set_title("주문 사전 4,869개", fontsize=15, pad=12, color="#2C2416")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    words = save_chart("learned_words.png", fig)

    fig, ax = plt.subplots(figsize=(7.2, 4.3))
    shops = ["국밥집", "카페", "햄버거집", "꽃집", "문구점", "선물가게"]
    first = [18736, 19033, 20073, 20454, 20867, 21897]
    rest = [14597, 14301, 13261, 12879, 12466, 11436]
    import numpy as np

    x = np.arange(len(shops))
    ax.bar(x, first, width=0.62, label="처음 읽어 준 메뉴", color="#C2510F")
    ax.bar(x, rest, width=0.62, bottom=first, label="그다음 메뉴들", color="#1F4E79")
    ax.set_xticks(x, shops, fontsize=11)
    ax.set_ylabel("주문한 사람", fontsize=12)
    ax.set_ylim(0, 42000)
    ax.set_title("가게마다 처음 메뉴가 많고, 나머지 메뉴도 주문됐다", fontsize=14, pad=12, color="#2C2416")
    ax.legend(frameon=False, ncol=2, loc="upper right")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    menus = save_chart("first_vs_rest.png", fig)
    return {"age": age, "why": why, "words": words, "menus": menus}


def paste_fit(canvas, path, box):
    image = Image.open(path).convert("RGB")
    x, y, bw, bh = box
    scale = min(bw / image.width, bh / image.height)
    nw, nh = max(1, int(image.width * scale)), max(1, int(image.height * scale))
    image = image.resize((nw, nh), Image.Resampling.LANCZOS)
    px, py = x + (bw - nw) // 2, y + (bh - nh) // 2
    canvas.paste(image, (px, py))
    return px, py, nw, nh


def card(draw, box, title, body, bar):
    rounded(draw, box, 18, CARD, LINE, 2)
    x, y, x2, y2 = box
    draw.rectangle((x, y, x + 10, y2), fill=bar)
    draw.text((x + 28, y + 16), title, font=font(26), fill=INK)
    draw_text(draw, (x + 28, y + 58), body, font(24), MUTED, max_w=(x2 - x - 48), leading=32)


TOTAL = 20
NOTES = []


def remember(image, note):
    NOTES.append((image, note))


def slide_title():
    image, draw = new_slide(rail=False)
    draw.rectangle((0, 0, 18, H), fill=ACCENT)
    draw.rectangle((0, 0, W, 18), fill=ACCENT)
    draw.text((80, 150), "개발 발표", font=font(28), fill=ACCENT)
    draw_text(
        draw,
        (80, 220),
        "초고령층을 위한\nSTT 기반 배리어프리\n인공지능 키오스크",
        font(68),
        INK,
        max_w=1500,
        leading=84,
    )
    draw.line((80, 520, 620, 520), fill=ACCENT, width=6)
    draw_text(
        draw,
        (80, 560),
        "햄버거집과 카페에서 직원에게 하시던 주문을\n키오스크가 짧은 말로 받게 만들었습니다.",
        font(32),
        MUTED,
        max_w=1500,
        leading=46,
    )
    chips = ["말로만 주문", "화면 확대", "소리 조절", "20만 명 가상 주문"]
    x = 80
    for chip in chips:
        face = font(24)
        tw = draw.textlength(chip, font=face)
        rounded(draw, (x, 760, x + tw + 36, 818), 18, CARD, LINE, 2)
        draw.text((x + 18, 774), chip, font=face, fill=NAVY)
        x += tw + 52
    draw.text((80, 980), "음성 인식은 기기 안의 Whisper 모델을 사용합니다.", font=font(22), fill=MUTED)
    remember(
        image,
        "안녕하세요. 초고령층이 말로 주문하는 배리어프리 키오스크를 발표하겠습니다. "
        "햄버거집과 카페에서 키오스크 대신 직원에게 주문하시던 장면을 보고 이 주제를 골랐습니다.",
    )


def slide_roadmap():
    image, draw = new_slide()
    header(draw, "오늘 이야기는 네 장면입니다")
    rows = [
        ("01", "왜 이 주제인가", "가게에서 본 장면과, 2024년에 발표된 서울 시민 조사"),
        ("02", "무엇을 만들었는가", "말로 받아 주는 주문, 확대, 소리, 4,869개 말"),
        ("03", "20만 명은 주문했는가", "청력, 시력, 이해력이 다른 사람이 키오스크 말만 따라 한 결과"),
        ("04", "기능이 서로 침범하는가", "같이 동작한 것과, 아직 손볼 네 곳"),
    ]
    y = 180
    for num, title, body in rows:
        rounded(draw, (72, y, 1848, y + 175), 22, CARD, LINE, 2)
        draw.text((100, y + 48), num, font=font(40), fill=ACCENT)
        draw.text((230, y + 36), title, font=font(34), fill=INK)
        draw.text((230, y + 96), body, font=font(26), fill=MUTED)
        y += 198
    footer(draw, 2, TOTAL)
    remember(image, "흐름만 먼저 말씀드리겠습니다. 이유, 만든 것, 20만 명 실험, 기능 간섭 순서입니다.")


def slide_reason():
    image, draw = new_slide()
    header(draw, "주제는 가게에서 본 장면에서 시작했습니다")
    blocks = [
        (ACCENT, "현장에서 본 것", "햄버거집과 카페처럼 키오스크가 자리 잡은 가게에서, 초고령층 손님이 화면 대신 직원을 찾아 주문하시는 장면을 반복해서 보았습니다."),
        (NAVY, "조사가 보여 준 것", "이 불편은 2021년 조사와 2023년 조사에 이어 남아 있습니다. 2024년 6월 발표에서 75세 이상의 키오스크 이용 경험은 19.1%입니다."),
        (TEAL, "그래서 만든 것", "그 가게의 음성 주문은 아직 자리에 없습니다. 직원이 짧게 받아 주듯, 키오스크가 듣고 짧게 되묻는 주문을 우리가 만들었습니다."),
    ]
    y = 170
    for color, title, body in blocks:
        rounded(draw, (72, y, 1848, y + 240), 22, CARD, LINE, 2)
        draw.rectangle((72, y, 86, y + 240), fill=color)
        draw.text((120, y + 28), title, font=font(32), fill=INK)
        draw_text(draw, (120, y + 90), body, font(28), MUTED, max_w=1660, leading=40)
        y += 262
    footer(draw, 3, TOTAL)
    remember(
        image,
        "주제 선정 이유입니다. 햄버거집과 카페에서 초고령층이 직원에게 주문하는 장면을 계속 보았습니다. "
        "2021년과 2023년 조사에도 이 어려움이 남아 있고, 음성으로 받아 주는 키오스크는 아직 그 가게에 없습니다. "
        "그래서 개발했습니다. 5,500명 조사는 서울디지털재단의 조사입니다. 우리가 그 설문을 한 것은 아닙니다.",
    )


def slide_age(charts):
    image, draw = new_slide()
    header(draw, "75세 이상은 키오스크를 써 본 비율이 19.1%입니다")
    paste_fit(image, charts["age"], (70, 160, 1180, 800))
    rounded(draw, (1280, 200, 1848, 900), 22, CARD, LINE, 2)
    draw.text((1320, 240), "조사", font=font(28), fill=ACCENT)
    draw_text(
        draw,
        (1320, 300),
        "서울디지털재단\n2023 서울시민\n디지털역량 실태조사\n\n만 19세 이상 5,500명\n고령층 2,500명 포함\n2024년 6월 12일 발표",
        font(26),
        INK,
        max_w=480,
        leading=38,
    )
    draw_text(
        draw,
        (1320, 680),
        "만 55세 이상 전체의\n이용 경험은 57.1%입니다.\n2년 전보다 11.3%포인트\n올랐고, 75세 이상은\n여전히 낮습니다.",
        font(24),
        MUTED,
        max_w=480,
        leading=34,
    )
    footer(draw, 4, TOTAL)
    remember(
        image,
        "왼쪽 그래프입니다. 55세에서 64세는 79.1%, 65세에서 74세는 50.4%, 75세 이상은 19.1%가 키오스크를 써 봤습니다. "
        "초고령층으로 갈수록 빈칸이 큽니다. 출처는 서울디지털재단의 2023년 조사, 2024년 6월 발표입니다.",
    )


def slide_why(charts):
    image, draw = new_slide()
    header(draw, "어려움의 이유는 속도, 선택, 용어, 글씨입니다")
    paste_fit(image, charts["why"], (40, 150, 1100, 720))
    items = [
        ("59.6%", "고령층이 이용 중에\n어려움을 겪음"),
        ("53.6%", "60대 이상은\n조작이 가장 불편"),
        ("23.2%", "60대 이상은\n작은 글씨가 불편"),
    ]
    y = 170
    for big, small in items:
        rounded(draw, (1160, y, 1848, y + 210), 20, CARD, LINE, 2)
        draw.text((1200, y + 36), big, font=font(48), fill=ACCENT)
        draw_text(draw, (1460, y + 48), small, font(26), INK, max_w=350, leading=36)
        y += 230
    draw.text((72, 900), "왼쪽: 서울디지털재단 2023 조사, 어려움을 겪은 고령층의 이유, 복수 응답.    오른쪽 아래 두 칸: 한국소비자원 2024 키오스크 이용 실태, 60대 이상.", font=font(18), fill=MUTED)
    footer(draw, 5, TOTAL)
    remember(
        image,
        "고령층의 59.6%가 키오스크를 쓰며 어려움을 겪었습니다. 이유는 뒷사람 눈치 53.6%, 선택 사항 46.3%, 용어 34.0%입니다. "
        "한국소비자원 2024년 조사에서는 60대 이상의 53.6%가 조작을, 23.2%가 작은 글씨를 불편하다고 했습니다. "
        "그래서 우리는 한 번에 하나만 짧게 묻고, 메뉴 이름을 읽어 주고, 글씨를 키우게 만들었습니다.",
    )


def slide_principles():
    image, draw = new_slide()
    header(draw, "직원이 받아 주듯, 네 가지를 지켰습니다")
    cells = [
        ("한 번에 하나", "긴 설명 대신 한 문장만 말합니다. 뒷사람 눈치를 줄이려면 질문이 짧아야 합니다."),
        ("메뉴를 추측하지 않음", "짧은 말이 두 메뉴에 걸치면 마음대로 고르지 않습니다. 다시 메뉴 이름을 묻습니다."),
        ("보이면 고른다", "글씨가 안 보이면 화면을 키웁니다. 주문을 추측해서 대신 넣지 않습니다."),
        ("들리면 따라 한다", "키오스크가 읽어 준 메뉴 이름을 손님이 따라 말하면 그 메뉴로 진행합니다."),
    ]
    positions = [(72, 180), (980, 180), (72, 560), (980, 560)]
    for (title, body), (x, y) in zip(cells, positions):
        card(draw, (x, y, x + 860, y + 330), title, body, ACCENT)
    footer(draw, 6, TOTAL)
    remember(
        image,
        "설계 원칙입니다. 한 문장만 말하고, 애매한 짧은 말로 메뉴를 찍지 않습니다. "
        "글씨가 안 보이면 확대하고, 키오스크가 읽어 준 이름을 따라 말하면 그 메뉴로 갑니다.",
    )


def slide_system():
    image, draw = new_slide()
    header(draw, "말은 기기 안에서 글로 바뀌고, 단계가 주문을 진행합니다")
    boxes = [
        (80, "손님의 말"),
        (440, "Whisper\n음성 인식"),
        (800, "주문 단계"),
        (1160, "짧은 음성\n안내"),
        (1520, "장바구니"),
    ]
    for x, label in boxes:
        rounded(draw, (x, 220, x + 280, 420), 20, CARD, LINE, 2)
        draw_text(draw, (x + 28, 270), label, font(28), INK, max_w=224, leading=38)
    rights = [x + 280 for x, _label in boxes[:-1]]
    lefts = [x for x, _label in boxes[1:]]
    for right, left in zip(rights, lefts):
        mid = (right + left) // 2
        draw.line((right + 14, 320, mid - 8, 320), fill=ACCENT, width=5)
        draw.polygon([(mid - 8, 308), (mid + 16, 320), (mid - 8, 332)], fill=ACCENT)
    bands = [
        ("학습한 말 4,869개", "같은 뜻의 말을 주문에 쓰는 말로 바꿉니다. 날씨 인사는 넣지 않았습니다."),
        ("여섯 가게", "국밥집, 카페, 햄버거집, 꽃집, 문구점, 선물가게. 가게마다 파는 메뉴가 다릅니다."),
        ("화면과 소리", "확대는 최대 3단계, 소리는 따로 키우고 줄입니다. 둘은 서로 다른 기능입니다."),
    ]
    x = 72
    for title, body in bands:
        rounded(draw, (x, 500, x + 580, 960), 20, CARD, LINE, 2)
        draw.text((x + 28, 530), title, font=font(30), fill=NAVY)
        draw_text(draw, (x + 28, 600), body, font(26), MUTED, max_w=520, leading=38)
        x += 608
    footer(draw, 7, TOTAL)
    remember(
        image,
        "구조입니다. 손님 말은 기기 안의 Whisper가 글로 바꿉니다. 그 글이 주문 단계를 움직이고, 키오스크는 짧은 문장으로 다시 말합니다. "
        "사전은 4,869개입니다. 가게는 여섯 곳입니다. 확대와 소리는 따로 둡니다.",
    )


def slide_flow():
    image, draw = new_slide()
    header(draw, "주문은 이 순서로 말합니다")
    steps = [
        ("1", "주문할게요", "주문을 엽니다"),
        ("2", "메뉴 이름", "치즈버거, 아메리카노"),
        ("3", "수량", "한 개, 두 잔. 100개는 거절"),
        ("4", "온도·잔", "커피는 온도, 주스는 잔 크기"),
        ("5", "더 담을까요", "네는 메뉴로, 아니요는 다음으로"),
        ("6", "포장·매장", "들고 갈지, 여기서 먹을지"),
        ("7", "카드·현금", "요약에서 결제, 화면에서 수단"),
        ("8", "안내", "아래에 수단을 넣어 달라는 안내"),
    ]
    for i, (num, title, body) in enumerate(steps):
        col, row = i % 4, i // 4
        x, y = 72 + col * 460, 190 + row * 380
        rounded(draw, (x, y, x + 430, y + 340), 22, CARD, LINE, 2)
        draw.ellipse((x + 24, y + 24, x + 92, y + 92), fill=ACCENT)
        tw = draw.textlength(num, font=font(32))
        draw.text((x + 58 - tw / 2, y + 38), num, font=font(32), fill=(255, 252, 247))
        draw.text((x + 28, y + 120), title, font=font(32), fill=INK)
        draw_text(draw, (x + 28, y + 190), body, font(24), MUTED, max_w=370, leading=34)
    footer(draw, 8, TOTAL)
    remember(
        image,
        "시연도 이 순서입니다. 주문할게요, 메뉴 이름, 수량, 커피는 온도, 주스는 잔 수 다음에 잔 크기, "
        "더 담지 않으면 아니요, 포장 또는 매장, 결제 다음 카드나 현금. 마지막은 4초 동안 수단을 넣어 달라는 안내입니다. "
        "실제 카드 단말기와 영수증은 아직 연결하지 않았습니다.",
    )


def slide_words(charts):
    image, draw = new_slide()
    header(draw, "고객의 말 1,311개를 정리하고, 같은 뜻 3,595개를 더했습니다")
    paste_fit(image, charts["words"], (40, 160, 1080, 780))
    rounded(draw, (1140, 190, 1848, 930), 22, CARD, LINE, 2)
    facts = [
        ("1,311 → 1,274", "고객이 한 말은 같은 말을 한 번만 남겨 1,274개가 되었습니다."),
        ("3,595 / 3,595", "뒤에 더한 말은 학습 검사에서 모두 배운 뜻으로 실행되었습니다. 차단은 0입니다."),
        ("날씨 인사", "“안녕하세요 날씨 좋네요”는 주문과 무관해서 사전에 없습니다. 첫 화면에 남습니다."),
    ]
    y = 220
    for title, body in facts:
        draw.text((1180, y), title, font=font(28), fill=ACCENT)
        y = draw_text(draw, (1180, y + 46), body, font(24), INK, max_w=620, leading=34)
        y += 36
    footer(draw, 9, TOTAL)
    remember(
        image,
        "말은 두 층입니다. 고객이 실제로 한 1,311개를 겹치면 1,274개입니다. 같은 주문으로 이어 주는 말 3,595개를 더해 사전은 4,869개입니다. "
        "3,595개는 검사에서 3,595개 모두 배운 뜻으로 실행되었고 차단은 0이었습니다. 날씨 인사는 일부러 넣지 않았습니다.",
    )


def slide_experiment():
    image, draw = new_slide()
    header(draw, "20만 명은 주문법을 모르고, 키오스크 말만 따라 했습니다")
    stats = [
        ("200,000", "서로 다른 사람"),
        ("200,000", "주문 성공"),
        ("0", "실패"),
        ("53", "가게-메뉴 조합"),
    ]
    x = 72
    for big, small in stats:
        rounded(draw, (x, 175, x + 430, 390), 20, CARD, LINE, 2)
        draw.text((x + 28, 200), big, font=font(48), fill=ACCENT)
        draw.text((x + 28, 275), small, font=font(26), fill=INK)
        x += 452
    rows = [
        ("청력", "점수가 낮고 문장이 길면 그 말을 알아듣지 못합니다. 점수를 올리면 통과하게 만들지 않았습니다."),
        ("시력", "점수가 낮으면 “글씨가 안 보여요”라고 말하고, 확대가 되기 전에는 메뉴 이름을 따라 하지 않습니다."),
        ("이해력", "한 문장에 뜻이 두 개면, 이해력이 낮을 때 그 문장을 못 알아듣습니다. “이해가 안 돼요” 뒤에 짧은 이름을 따라 합니다."),
        ("성공", "장바구니에 이름 있는 메뉴가 1개 이상 담기고, 안내 또는 완료 화면까지 가면 성공입니다."),
    ]
    y = 430
    for title, body in rows:
        draw.text((90, y), title, font=font(28), fill=NAVY)
        draw_text(draw, (250, y), body, font(24), INK, max_w=1580, leading=32)
        y += 130
    footer(draw, 10, TOTAL)
    remember(
        image,
        "실험입니다. 20만 명은 모두 다른 사람입니다. 주문 방법을 미리 알지 못하고, 키오스크가 읽어 주는 말만 따라 합니다. "
        "청력, 시력, 이해력을 각각 봤습니다. 알아듣는 기준은 느슨하게 풀지 않았습니다. "
        "성공은 20만 명 모두, 실패는 0입니다. 53개 가게-메뉴 조합이 한 번 이상 주문됐습니다. "
        "이것은 조용한 가상 화면의 결과입니다. 매장 소음은 넣지 않았습니다.",
    )


def slide_figure(title, subtitle, path, page, note, caption=""):
    image, draw = new_slide()
    y = header(draw, title, subtitle)
    caption_h = 78 if caption else 0
    bottom = 990
    paste_fit(image, path, (48, y, 1824, bottom - caption_h - y))
    if caption:
        draw_text(draw, (72, bottom - caption_h + 6), caption, font(20), MUTED, max_w=1760, leading=28)
    footer(draw, page, TOTAL)
    remember(image, note)


def slide_check():
    image, draw = new_slide()
    header(draw, "기능을 이어서 말해 보니, 주문 줄기는 끝까지 갔습니다")
    chips = [
        "기본 주문", "확대 후 주문", "소리 키워도 주문", "크게 말은 확대 아님",
        "사용법 후 주문", "날씨는 메뉴 아님", "시작 전 결제 유지", "다음 메뉴",
        "수량 중 확대", "다섯 개는 5", "백 개 거절", "포장과 매장",
        "결제 화면의 현금", "카드 말고 현금", "주스 중간 잔", "이전으로",
        "가게 변경", "소리 줄여",
    ]
    x, y = 72, 175
    face = font(24)
    for chip in chips:
        tw = draw.textlength(chip, font=face)
        if x + tw + 40 > 1850:
            x = 72
            y += 78
        rounded(draw, (x, y, x + tw + 36, y + 62), 16, (232, 244, 239), (190, 214, 204), 2)
        draw.text((x + 18, y + 14), chip, font=face, fill=TEAL)
        x += tw + 52
    rounded(draw, (72, 620, 1848, 960), 22, CARD, LINE, 2)
    draw.text((104, 650), "같이 확인한 장면", font=font(28), fill=NAVY)
    draw_text(
        draw,
        (104, 710),
        "치즈버거 두 개 포장 카드, 국밥 다섯 개, 아메리카노 온도, 오렌지 주스 중간 잔까지 각 단계가 자기 일을 끝냈습니다.\n"
        "주문 전에 “결제”나 “다음”을 말하면 첫 화면에 남습니다. 결제 화면에서 “결제”를 반복해도 결제가 한 번 더 나가지 않습니다.\n"
        "요약에서 “카드 말고 현금”은 현금으로 처리됐습니다.",
        font(26),
        INK,
        max_w=1680,
        leading=40,
    )
    footer(draw, 15, TOTAL)
    remember(
        image,
        "기능 간섭 검사입니다. 주문, 확대, 소리, 사용법, 다음 메뉴, 수량, 포장, 결제를 이어서 말했습니다. "
        "주문 줄기는 서로 지우지 않고 끝까지 갔습니다. 카드 말고 현금은 현금으로 끝났습니다. "
        "이어서 아직 침범이 남은 곳을 말씀드리겠습니다.",
    )


def slide_gaps():
    image, draw = new_slide()
    header(draw, "아직 한 기능이 다른 기능을 가져가는 곳이 네 곳입니다")
    gaps = [
        ("1", "다음 손님", "처음 화면이 앞 손님의 확대, 소리, 메뉴 순서를 지우지 않습니다. 확대 1단계와 메뉴 위치 2가 그대로 남았습니다."),
        ("2", "포장·카드·현금", "수량을 물을 때도 “포장”이 장소를 바꿉니다. 메뉴를 더 물을 때 “카드”는 결제된 것처럼 말하고, 결제 수단은 비어 있습니다."),
        ("3", "수량에서 네", "다른 화면의 “네”는 동의입니다. 수량을 물을 때 “네”는 4개가 되어 다음 단계로 넘어갑니다."),
        ("4", "이해가 안 돼요", "세트 화면에서 “단품이요”라고 말하지만 단품으로 넘어가지 않습니다. 수량에서는 “한 개 해 주세요”라고만 말합니다."),
    ]
    for i, (num, title, body) in enumerate(gaps):
        col, row = i % 2, i // 2
        x, y = 72 + col * 920, 175 + row * 390
        rounded(draw, (x, y, x + 880, y + 360), 22, CARD, LINE, 2)
        draw.ellipse((x + 24, y + 24, x + 84, y + 84), fill=GOLD)
        draw.text((x + 44, y + 34), num, font=font(28), fill=CARD)
        draw.text((x + 108, y + 36), title, font=font(30), fill=INK)
        draw_text(draw, (x + 28, y + 120), body, font(24), MUTED, max_w=820, leading=34)
    footer(draw, 16, TOTAL)
    remember(
        image,
        "침범이 남은 네 곳입니다. 첫째, 처음 화면이 앞 손님의 확대와 메뉴 순서를 남깁니다. 소리도 처음 화면 코드가 되돌리지 않습니다. "
        "둘째, 포장과 카드, 현금이 자기 단계가 아닐 때도 그 말을 가져갑니다. 카드는 결제된 것처럼 말하고 결제 수단은 비어 있습니다. "
        "셋째, 수량에서 네는 4개입니다. 넷째, 이해가 안 돼요는 답을 읽어 주기만 하고 그 선택을 확정하지 않습니다. "
        "학습 별칭이 4,549개 쌓인 브라우저에서는 다음이 네로 바뀌어 수량 4개가 되었고, 더 크게 말해 줘가 소리를 키우지 못했습니다. "
        "그 기록을 지우자 두 말은 제 기능으로 돌아왔습니다.",
    )


def slide_demo():
    image, draw = new_slide()
    header(draw, "발표 중 시연은 이 여덟 마디면 됩니다")
    lines = ["주문할게요", "글씨가 안 보여요", "치즈버거", "단품", "한 개", "아니요", "포장", "카드"]
    x = 72
    for i, line in enumerate(lines):
        rounded(draw, (x, 190, x + 210, 360), 18, CARD, LINE, 2)
        draw.text((x + 18, 210), f"{i + 1}", font=font(22), fill=ACCENT)
        face = font(26 if len(line) < 8 else 22)
        draw_text(draw, (x + 16, 255), line, face, INK, max_w=178, leading=32)
        x += 228
    rounded(draw, (72, 410, 1848, 960), 22, CARD, LINE, 2)
    draw.text((110, 450), "옆에 한 번씩만 더 보여 줄 말", font=font(30), fill=NAVY)
    extras = [
        ("소리 키워", "목소리가 커지고, 글씨 크기는 그대로입니다."),
        ("사용법 알려줘", "주문은 메뉴 이름으로 한다는 짧은 안내가 나옵니다."),
        ("메뉴판 보여줘, 다음", "치즈버거 다음 불고기버거를 읽어 줍니다."),
        ("다섯 개", "국밥 수량은 5입니다. 백 개는 받지 않습니다."),
    ]
    y = 520
    for title, body in extras:
        draw.text((110, y), title, font=font(28), fill=ACCENT)
        draw.text((520, y), body, font=font(26), fill=INK)
        y += 90
    footer(draw, 17, TOTAL)
    remember(
        image,
        "시연은 여덟 마디입니다. 주문할게요, 글씨가 안 보여요, 치즈버거, 단품, 한 개, 아니요, 포장, 카드. "
        "시간이 있으면 소리 키워, 사용법, 메뉴판에서 다음, 다섯 개를 하나씩 보여 주십시오. "
        "화면이 안 바뀌면 Ctrl+F5입니다. 파일은 폴더째 열어야 합니다.",
    )


def slide_close():
    image, draw = new_slide()
    header(draw, "키오스크가 읽어 준 말로, 20만 명의 주문이 끝났습니다")
    points = [
        "75세 이상의 키오스크 이용 경험 19.1%에서 출발했습니다.",
        "말은 기기 안에서 인식하고, 4,869개 사전과 짧은 문장으로 받습니다.",
        "청력, 시력, 이해력이 다른 20만 명 모두 키오스크 안내만으로 주문에 성공했습니다.",
        "53개 메뉴가 한 번 이상 주문됐고, 가장 적은 메뉴도 809명이 주문했습니다.",
        "처음 화면의 확대 유지, 수량에서 네, 포장·결제 말, 이해 안내는 다음 수정입니다.",
    ]
    y = 180
    for point in points:
        draw.ellipse((84, y + 10, 108, y + 34), fill=ACCENT)
        y = draw_text(draw, (130, y), point, font(30), INK, max_w=1680, leading=42)
        y += 28
    rounded(draw, (72, 780, 1848, 960), 22, (255, 244, 232), (232, 196, 166), 2)
    draw_text(
        draw,
        (104, 820),
        "가상 화면의 성공을 매장 성공으로 말하지 않겠습니다.\n소음 속의 다양한 말투와 실제 결제 단말기는 다음 과제입니다.",
        font(28),
        INK,
        max_w=1680,
        leading=42,
    )
    footer(draw, 18, TOTAL)
    remember(
        image,
        "정리입니다. 초고령층이 직원에게 하시던 주문을, 키오스크가 짧게 듣고 짧게 되묻게 만들었습니다. "
        "20만 명은 키오스크가 읽어 준 말로 주문에 성공했습니다. 이 성공은 조용한 가상 화면의 결과입니다. "
        "다음 손님에게 확대가 남지 않게, 수량에서 네가 4개가 되지 않게, 포장과 결제 말이 자기 화면에서만 동작하게 고치겠습니다. 들어 주셔서 감사합니다.",
    )


def build():
    charts = make_charts()
    slide_title()
    slide_roadmap()
    slide_reason()
    slide_age(charts)
    slide_why(charts)
    slide_principles()
    slide_system()
    slide_flow()
    slide_words(charts)
    slide_experiment()
    slide_figure(
        "20만 명의 청력, 시력, 이해력",
        "점선은 도움을 나눈 기준입니다. 청력 35 미만, 확대가 필요한 시력 45 미만, 한 가지만 이해하는 이해력 55 미만.",
        PAPER / "paper_fig1_profiles.png",
        11,
        "능력 분포입니다. 청력, 시력, 이해력은 사람마다 다르게 퍼져 있습니다. "
        "점선은 청력 35, 시력 45, 이해력 55입니다. 이 기준 아래도 주문에서 빼지 않았습니다.",
    )
    slide_figure(
        "확대가 필요한 사람도 주문에 성공했습니다",
        "오른쪽 세 집단은 전체 20만 명 안에 있고 서로 겹칩니다. 막대를 더해 인원을 세지 않습니다.",
        PAPER / "paper_fig2_groups.png",
        12,
        "확대가 필요 없는 사람 80,788명, 1단계 48,422명, 2단계 38,577명, 3단계 32,213명입니다. "
        "청력 35 미만 52,265명, 이해력 55 미만 106,070명, 확대가 필요한 사람 119,212명입니다. "
        "세 집단은 겹칩니다. 모두 성공했습니다.",
    )
    slide_figure(
        "여섯 가게에 거의 같은 수가 들어갔습니다",
        "국밥집 33,333명, 카페와 햄버거집 33,334명, 꽃집·문구점·선물가게 33,333명.",
        PAPER / "paper_fig5_shops.png",
        13,
        "가게 인원입니다. 여섯 가게가 각각 약 3만 3천 명입니다. 한 가게에 몰리지 않았습니다.",
    )
    slide_figure(
        "처음 읽어 준 메뉴가 많고, 나머지 메뉴도 비지 않았습니다",
        "이해력이 낮으면 “이름, 다음”에서 다음을 못 알아듣고 지금 이름을 주문합니다. 나머지 최소는 국밥집 오렌지 주스 809명입니다.",
        charts["menus"],
        14,
        "메뉴 분포입니다. 각 가게의 첫 메뉴가 큰 이유는, 이해력이 낮으면 다음이라는 말을 못 알아듣고 지금 들리는 이름을 주문하기 때문입니다. "
        "그래도 53개 조합이 모두 주문됐고, 가장 적은 국밥집 오렌지 주스는 809명입니다. "
        "국밥집 첫 메뉴를 뺀 나머지 14,597명, 카페 14,301명, 햄버거집 13,261명, 꽃집 12,879명, 문구점 12,466명, 선물가게 11,436명입니다.",
        "처음 메뉴: 돼지국밥 18,736 · 아메리카노 19,033 · 치즈버거 20,073 · 장미 한 송이 20,454 · 볼펜 20,867 · 손수건 21,897",
    )
    slide_check()
    slide_gaps()
    slide_demo()
    slide_close()

    # 논문용 전체 메뉴 그림은 발표 본편 뒤에 참고 슬라이드로 둔다.
    extra_notes = []
    extras = [
        (
            "참고. 가게별 전체 메뉴",
            "붉은 막대는 그 가게가 처음 읽어 준 메뉴입니다. 발표에서는 앞의 요약 그림을 말하면 됩니다.",
            PAPER / "paper_fig3_menus.png",
            "참고 그림입니다. 전체 메뉴입니다. 발표 시간에는 앞 장의 요약만 말씀하시면 됩니다.",
        ),
        (
            "참고. 처음 메뉴를 뺀 나머지",
            "막대 길이를 나머지끼리 비교한 그림입니다. 최소 809명, 음식점 나머지는 약 1천 명에서 3천 명대입니다.",
            PAPER / "paper_fig4_other_menus.png",
            "참고 그림입니다. 처음 메뉴를 뺀 나머지입니다. 국밥집 나머지가 800명대, 다른 가게는 그보다 많습니다.",
        ),
    ]
    for offset, (title, subtitle, path, note) in enumerate(extras, start=19):
        image, draw = new_slide()
        y = header(draw, title, subtitle)
        paste_fit(image, path, (48, y, 1824, 990 - y))
        footer(draw, offset, TOTAL)
        extra_notes.append((image, note))

    # 본편 쪽수를 이미 18로 찍었으므로 참고 쪽수는 19, 20으로 다시 그린다.
    NOTES.extend(extra_notes)

    SLIDES.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, (image, _note) in enumerate(NOTES, start=1):
        path = SLIDES / f"slide_{i:02d}.png"
        image.save(path, "PNG")
        paths.append(path)

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "초고령층을 위한 STT 기반 배리어프리 인공지능 키오스크"
    prs.core_properties.subject = "말로 주문하는 배리어프리 키오스크 개발 발표"
    blank = prs.slide_layouts[6]
    for path, (_image, note) in zip(paths, NOTES):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        notes = slide.notes_slide.notes_text_frame
        notes.text = note
    prs.save(OUT)
    print(f"slides {len(paths)} -> {OUT}")


if __name__ == "__main__":
    build()
